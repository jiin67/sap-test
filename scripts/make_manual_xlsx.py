# make_manual_xlsx.py : build the Excel version of the program manual
import io, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

BASE = r"C:\sap-test\results"
SHOTS = os.path.join(BASE, "shots")
THUMB = os.path.join(BASE, "shots", "_man")
OUT = os.path.join(BASE, "ZCRM052_사용설명서_20260910.xlsx")
os.makedirs(THUMB, exist_ok=True)

F = "Arial"
MONO = "Consolas"
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, bold=True, size=10)
TITLE = Font(name=F, bold=True, size=15)
SEC = Font(name=F, bold=True, size=12, color="1F3864")
NOTE = Font(name=F, size=9, italic=True, color="666666")
WARN = Font(name=F, size=10, bold=True, color="C00000")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
NG_FILL = PatternFill("solid", fgColor="FCE4D6")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
IN_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def thumb(fn, maxw=820):
    src = os.path.join(SHOTS, fn)
    if not os.path.exists(src):
        return None
    dst = os.path.join(THUMB, fn)
    im = PILImage.open(src)
    im.thumbnail((maxw, maxw))
    im.save(dst, optimize=True)
    return dst


def hdr(ws, row, cells):
    for j, v in enumerate(cells, start=1):
        c = ws.cell(row=row, column=j, value=v)
        c.fill = H_FILL
        c.font = H_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    return row + 1


def row(ws, row_i, cells, font=BODY, fill=None):
    for j, v in enumerate(cells, start=1):
        c = ws.cell(row=row_i, column=j, value=v)
        c.font = font
        c.border = BOX
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if fill:
            c.fill = fill
    return row_i + 1


def section(ws, row_i, text):
    c = ws.cell(row=row_i, column=1, value=text)
    c.font = SEC
    return row_i + 1


def widths(ws, spec):
    for k, v in spec.items():
        ws.column_dimensions[k].width = v


def put_image(ws, row_i, fn, caption=None, col="A"):
    p = thumb(fn)
    if not p:
        return row_i
    if caption:
        c = ws.cell(row=row_i, column=1, value=caption)
        c.font = BOLD
        row_i += 1
    img = XLImage(p)
    ws.add_image(img, "{}{}".format(col, row_i))
    return row_i + int(img.height / 19) + 2


wb = Workbook()

# ============================================================ 1. 개요와 흐름
ws = wb.active
ws.title = "1.개요와 흐름"
ws["A1"] = "회수품 적합성 관리 사용 설명서"
ws["A1"].font = TITLE
ws["A2"] = "T-code ZCRM052 (일반) / ZCRM052S (구독 전용)"
ws["A2"].font = NOTE
r = 4
r = hdr(ws, r, ["항목", "내용"])
for k, v in [
    ("대상 프로그램", "ZCRM052 → ZCRMR052 (일반)  /  ZCRM052S → ZCRMR052S (구독 전용)"),
    ("작성 기준", "KDQ 클라이언트 210, 로그온 언어 KO"),
    ("작성일", "2026-09-10"),
    ("두 프로그램 관계", "구독 데이터 필터를 빼면 화면과 동작이 같다"),
]:
    r = row(ws, r, [k, v])
r += 1
r = row(ws, r, ["표기", "의미"], H_FONT, H_FILL)
r = row(ws, r, ["[실측]", "이 문서 작성 과정에서 화면·DB 로 직접 확인한 내용"], BODY, OK_FILL)
r = row(ws, r, ["[개발가이드]", "개발 문서에 근거한 내용. 화면으로 직접 확인하지 않았다"], BODY, WARN_FILL)
r += 1

r = section(ws, r, "전체 흐름")
flow = [
    "① 선택화면 (screen 1000)",
    "     조회조건 입력  →  실행 (F8)",
    "          │",
    "          │   ZCRM_CICSVC ─┬─ ZCRM_REQ ─┬─ ZCRM_RECPT",
    "          │   (상담·서비스)  │ (회수품요청) │ (회수품접수)",
    "          │                 │            └─ ZCRM_RECPT_WEAR",
    "          │                 └─ MVKE (제품계층, INNER JOIN)",
    "          ▼",
    "② 결과화면 (screen 100)",
    "     ALV 목록  +  접수부서 / 접수자 입력란",
    "          │",
    "          │   적합수량 · 부적합수량 · 부적합사유 · 책임사업장 · 입고일자 입력",
    "          ▼   저장 (Ctrl+S)  →  \"저장 하시겠습니까?\"  →  예",
    "③ 업무 검증 6종 통과 시 저장",
    "          │",
    "          ├─ ZCRM_RECPTNO      접수번호 채번 (일자 + 순번)",
    "          ├─ ZCRM_RECPT        접수 · 적합성 판정 결과",
    "          └─ ZCRM_RECPT_WEAR   회수품 실제 입고일자",
]
for line in flow:
    c = ws.cell(row=r, column=1, value=line)
    c.font = Font(name=MONO, size=10)
    r += 1
r += 1
r = section(ws, r, "시트 안내")
r = hdr(ws, r, ["시트", "내용"])
for k, v in [
    ("2.선택화면", "조회조건 항목과 라디오에 따른 화면 변화"),
    ("3.조회 테이블", "실행 시 읽어오는 테이블과 연결조건"),
    ("4.결과화면", "화면 구성, 버튼, ALV 컬럼 45개"),
    ("5.입력과 저장", "입력 순서, 검증 6종, 저장되는 테이블"),
    ("6.상세조회와 종료", "회수가능 제조일 상세, 화면 종료"),
    ("7.자주 겪는 상황", "증상별 원인과 조치"),
    ("8.두 프로그램 차이", "일반 / 구독 전용 비교"),
    ("9.화면 모음", "캡처 모음"),
]:
    r = row(ws, r, [k, v])
widths(ws, {"A": 26, "B": 92})

# ============================================================ 2. 선택화면
ws = wb.create_sheet("2.선택화면")
ws["A1"] = "선택화면 (screen 1000)"
ws["A1"].font = TITLE
r = 3
r = section(ws, r, "조회조건 항목")
r = hdr(ws, r, ["순서", "항목", "기술명", "필수", "설명"])
sel = [
    ("1", "처리조직", "PA_R01 / PA_R02", "필수", "라디오. 센터조치 / 서비스점조치"),
    ("2", "센터", "SO_VKBUR", "선택", "센터조치일 때만 표시"),
    ("2'", "서비스점", "SO_KUNNR", "선택", "서비스점조치일 때만 표시"),
    ("3", "회수품요청일자", "SO_PRDRE", "선택", "청구요청일 범위"),
    ("4", "상담번호", "SO_CSLTN", "선택", "범위 지정 가능"),
    ("5", "접수일자", "SO_RCDT", "선택", "회수품 접수일 때만 표시. 기본값 오늘~오늘"),
    ("6", "제품 계층구조", "SO_PRDH", "선택", "MVKE-PRODH 기준"),
    ("7", "유/무상", "SO_PRTY", "선택", "단일값. Y 유상 / N 무상"),
    ("8", "자재번호(부품)", "SO_PRT", "선택", ""),
    ("9", "멀티킷 사용 여부", "SO_MULTK", "선택", "단일값. 화면 O / X"),
    ("10", "접수여부", "PA_CHX / PA_NCH", "필수", "라디오. 회수품 접수 / 회수품 미접수"),
]
for s in sel:
    r = row(ws, r, list(s))
r += 1
r = section(ws, r, "[실측] 라디오에 따라 조회조건이 나타나고 사라진다")
ws.cell(row=r, column=1, value=(
    "조회조건은 회색으로 비활성화되는 것이 아니라 화면에서 아예 사라진다. "
    "라디오를 바꾸면 즉시 반영되고, 되돌리면 원래 구성으로 정확히 복원된다.")).font = BODY
r += 2
r = hdr(ws, r, ["선택", "나타나는 조건", "사라지는 조건"])
for a, b, c2 in [
    ("처리조직 = 센터조치", "센터", "서비스점"),
    ("처리조직 = 서비스점조치", "서비스점", "센터"),
    ("접수여부 = 회수품 접수", "접수일자", "-"),
    ("접수여부 = 회수품 미접수", "-", "접수일자"),
]:
    r = row(ws, r, [a, b, c2], BODY, OK_FILL)
r += 1
r = section(ws, r, "[실측] 멀티킷 입력값 변환")
r = hdr(ws, r, ["화면 입력", "DB ZCRM_CICSVC-MULTIKITYN"])
r = row(ws, r, ["O", "X (멀티킷 사용)"])
r = row(ws, r, ["X", "공백 (멀티킷 미사용)"])
r += 1
r = section(ws, r, "실행")
ws.cell(row=r, column=1, value="실행 버튼 또는 F8").font = BODY
r += 2
ws.cell(row=r, column=1, value=(
    "[실측 주의] 조회 범위가 넓으면 매우 오래 걸린다. 센터조치 + 회수품 접수 기준으로 "
    "2026.01.01~09.09 를 조회하면 71,912건이고 수 분이 걸린다. 15일 이내로 끊어서 "
    "조회할 것을 권장한다. 15일이면 센터조치 약 5,300건, 서비스점조치 약 1,000건이다.")).font = WARN
r += 3
r = put_image(ws, r, "MANUAL_sel_center.png", "센터조치 선택 - 센터 조회조건이 보인다")
r = put_image(ws, r, "MANUAL_sel_dealer.png", "서비스점조치 선택 - 서비스점 조회조건으로 바뀐다")
r = put_image(ws, r, "MANUAL_sel_notrecv.png", "회수품 미접수 선택 - 접수일자가 사라진다")
widths(ws, {"A": 22, "B": 22, "C": 22, "D": 10, "E": 52})

# ============================================================ 3. 조회 테이블
ws = wb.create_sheet("3.조회 테이블")
ws["A1"] = "실행 시 읽어오는 테이블"
ws["A1"].font = TITLE
r = 3
r = section(ws, r, "[개발가이드] 조회 흐름과 연결조건")
r = hdr(ws, r, ["순서", "테이블", "내용", "Key", "상위와의 연결조건"])
flowt = [
    ("1", "ZCRM_CICSVC", "상담·서비스 처리 원천", "CSLTNO + CSLTNO_ITEM", "-"),
    ("2", "ZCRM_REQ", "회수품 요청", "REQNO + REQNO_ITEM",
     "CSLTNO = CSLTNO, CSLTNO_ITEM = CSLTNO_ITEM"),
    ("3", "ZCRM_RECPT", "회수품 접수·적합성 판정", "RECPTNO + RECPTNO_ITEM",
     "REQNO = REQNO, REQNO_ITEM = REQNO_ITEM"),
    ("4", "ZCRM_RECPT_WEAR", "회수품 실제 입고일자", "RECPTNO + RECPTNO_ITEM",
     "RECPTNO = RECPTNO, RECPTNO_ITEM = RECPTNO_ITEM"),
    ("5", "ZCRM_QRST_PRT", "품질검사 결과·마감상태", "RECPTNO + RECPTNO_ITEM",
     "RECPTNO = RECPTNO, RECPTNO_ITEM = RECPTNO_ITEM"),
    ("*", "MVKE", "제품 계층 (INNER JOIN)", "MATNR + VKORG + VTWEG",
     "MATNR = ZCRM_CICSVC-MATNR, VKORG = ASKR, VTWEG = 10"),
]
for t in flowt:
    r = row(ws, r, list(t))
r += 1
ws.cell(row=r, column=1, value=(
    "MVKE 에 해당 자재의 ASKR / 10 조합이 없으면 INNER JOIN 대상에서 빠져 조회되지 않는다.")).font = WARN
r += 2
r = section(ws, r, "[개발가이드] 접수 / 미접수 구분")
r = hdr(ws, r, ["접수여부", "대상"])
r = row(ws, r, ["회수품 접수", "ZCRM_RECPT 에 해당 요청이 있고 STATUS = 'X' 인 건"])
r = row(ws, r, ["회수품 미접수", "ZCRM_RECPT 에 없거나 STATUS <> 'X' 인 건"])
r += 1
ws.cell(row=r, column=1, value=(
    "공통 전제조건: ZCRM_REQ-TAGNO1 이 있어야 하고 ZCRM_REQ-REQ_DROP <> 'X' 여야 한다.")).font = BODY
r += 2
r = section(ws, r, "구독 전용 프로그램의 추가 조건")
c = ws.cell(row=r, column=1, value="AND zcrm_cicsvc~subs_type = gc_subs_type_s    \" SUBS_TYPE = 'S'")
c.font = Font(name=MONO, size=10, bold=True, color="C00000")
r += 1
ws.cell(row=r, column=1, value=(
    "ZCRMR052S 는 접수 SELECT 와 미접수 SELECT 양쪽에 위 조건이 붙는다. "
    "[실측] 2026-09-10 현재 비교 테스트를 위해 주석 처리되어 있어 일반 데이터도 조회된다. "
    "테스트 종료 후 반드시 원복해야 한다.")).font = WARN
r += 3
r = section(ws, r, "[개발가이드] 화면 표시용으로 함께 읽는 데이터")
r = hdr(ws, r, ["화면 항목", "원천 테이블과 조건"])
for a, b in [
    ("대체자재번호·대체사유", "KOTD002 (KSCHL='Z001', VKORG='ASKR', 유효기간) → KONDD"),
    ("대체품 자재명·구품번", "MAKT, MARA-BISMT, 대체사유명 TVSUT"),
    ("회수가능 제조일·보증기간", "MVKE-PRODH → ZCRM_WRTGROUP (WRT_TYPE='A')"),
    ("단납 보증기간", "ZCRM_CICSVC-SERNR → ZCRM_WRTDG (WRTPER<>0, LOEVM<>'X'). "
                  "CIC_YN='C' 이면 보증기간에 9개월 가산"),
    ("해지 서비스점 자동 부적합", "KNVV-AUFSD (ASKR/10/10) = 'Z8', 공백이면 KNA1-AUFSD = 'Z8'. "
                          "해당 시 NCF='N', NCF_RSN_V='Z10' 자동 설정"),
    ("해피콜 부적합", "HNCFRCD 가 지정 6개 코드이고 RSLTDATE >= 20140701 일 때만 표시"),
    ("적합 금액", "A911 (VKORG='ASKR', VTWEG='10', KSCHL='PR00', 유효기간) → KONP"),
]:
    r = row(ws, r, [a, b])
widths(ws, {"A": 24, "B": 20, "C": 26, "D": 26, "E": 52})

# ============================================================ 4. 결과화면
ws = wb.create_sheet("4.결과화면")
ws["A1"] = "결과화면 (screen 100)"
ws["A1"].font = TITLE
r = 3
r = section(ws, r, "화면 구성")
r = hdr(ws, r, ["영역", "기술명", "설명"])
for a, b, c2 in [
    ("창 제목", "-", "회수품적합성관리 (건수). 신규는 [구독전용] 회수품 적합성 관리 (건수)"),
    ("접수부서명", "GS_HEADER-REMARK1", "저장 시 필수. 음영 표시"),
    ("접수자", "GS_HEADER-REMARK2", "저장 시 필수"),
    ("ALV 목록", "cntlCON1", "조회 결과"),
]:
    r = row(ws, r, [a, b, c2])
r += 1
ws.cell(row=r, column=1, value="[실측] 조회 건수는 창 제목의 괄호 안 숫자로 확인한다.").font = BODY
r += 2
r = section(ws, r, "상단 버튼")
r = hdr(ws, r, ["버튼", "단축키", "동작"])
r = row(ws, r, ["회수가능 제조일 상세", "Shift+F11", "선택한 1행의 보증기간 상세 팝업"], BODY, OK_FILL)
r = row(ws, r, ["부적합 사유", "Ctrl+F2", "[실측] 동작하지 않음. '유효한 기능을 선택하십시오.'"], BODY, NG_FILL)
r = row(ws, r, ["책임사업장", "Ctrl+F3", "[실측] 동작하지 않음. '유효한 기능을 선택하십시오.'"], BODY, NG_FILL)
r = row(ws, r, ["저장", "Ctrl+S", "입력값 검증 후 저장"], BODY, OK_FILL)
r = row(ws, r, ["뒤로", "F3", "선택화면으로 복귀"], BODY, OK_FILL)
r += 1
ws.cell(row=r, column=1, value=(
    "부적합 사유·책임사업장 버튼은 두 프로그램 모두 반응이 없다. 행 선택 유무나 셀 위치를 "
    "바꿔도 같다. 값은 ALV 셀에서 직접 입력하면 되므로 업무에는 지장이 없다.")).font = BODY
r += 3
r = section(ws, r, "ALV 컬럼 (전체 45개)")
ws.cell(row=r, column=1, value=(
    "실제 표시 여부는 ALV 레이아웃(변형)에 따라 달라진다. 입력 표시가 있는 6개가 사용자가 "
    "값을 넣는 컬럼이다.")).font = BODY
r += 2
r = hdr(ws, r, ["기술명", "컬럼명", "표시", "입력"])
cols = [
    ("CSLTNO", "상담번호", "공통", ""), ("REQNO", "요청번호", "공통", ""),
    ("RECPTNO", "접수번호", "공통", ""), ("RECPTDATE", "접수일자", "공통", ""),
    ("SRV_QTY", "요청수량", "공통", ""),
    ("CF_QTY", "적합수량", "공통", "입력"),
    ("NCF_QTY", "부적합수량", "공통", "입력"),
    ("NCFT", "적합구분", "공통", ""),
    ("NCF_RSN_VT", "부적합사유", "공통", "입력(드롭다운)"),
    ("NCF_RSN_V", "부적합사유(코드)", "기존만", ""),
    ("RSPC", "책임사업장", "공통", "입력"),
    ("RSPCT", "책임사업장명", "공통", ""),
    ("REMARK", "비고", "공통", "입력"),
    ("WEARDATS", "회수품입고일자", "공통", "입력"),
    ("REMARK1", "접수부서", "공통", ""), ("REMARK2", "접수자명", "공통", ""),
    ("MATNR_PRT", "자재번호", "공통", ""), ("MAKTX_PRT", "자재내역", "공통", ""),
    ("MATNR_PRT_OLD", "구자재번호(회수품)", "공통", ""),
    ("MATNR_BASE", "기본자재내역", "공통", ""),
    ("SMATN", "대체자재번호", "공통", ""), ("SMATNT", "대체자재내역", "공통", ""),
    ("SMATN_OLD", "구자재번호(대체품)", "공통", ""),
    ("SUGRD", "대체사유코드", "공통", ""), ("SUGRDT", "대체사유명", "공통", ""),
    ("TAGNO1", "꼬리표NO", "공통", ""), ("TAGNO1_CHECKT", "불일치유무", "공통", ""),
    ("PRDDATE_PRT", "부품제조일", "공통", ""), ("PRDRETDATE", "회수가능제조일", "공통", ""),
    ("PRDREQDATE", "청구요청일", "공통", ""), ("RSLTDATE", "서비스 조치일자", "공통", ""),
    ("CONFDAT", "적합판정일", "공통", ""),
    ("PRICETYPET", "서비스 처리결과(유상/무상)", "공통", ""),
    ("REQTYPE_T", "청구유형", "공통", ""), ("CSYMP_T2", "현장증상2", "공통", ""),
    ("HNCFRCD", "H/C부적합사유", "공통", ""), ("HNCFRSN", "H/C부적합사유명", "공통", ""),
    ("KUNNR_SVC", "서비스점코드", "공통", ""), ("SVC_NAME2", "서비스점", "공통", ""),
    ("VKBUR", "센터코드", "공통", ""), ("VKBURT", "센터명", "공통", ""),
    ("MEINS", "기본 단위", "공통", ""),
    ("STATUST", "상태", "신규만", ""),
    ("MULTIKITYN", "멀티킷 사용여부", "신규만", ""),
    ("SUBS_TYPE", "구독구분", "신규만", ""),
]
for a, b, c2, d in cols:
    fill = IN_FILL if d else (WARN_FILL if c2 != "공통" else None)
    r = row(ws, r, [a, b, c2, d], BODY, fill)
r += 1
ws.cell(row=r, column=1, value=(
    "[실측] 기존 프로그램에는 저장된 ALV 레이아웃이 6개 있고 /CRM01 이 기본 설정이라 42개 "
    "컬럼만 그 순서로 보인다. 신규 프로그램에는 레이아웃이 없어 원본 44개가 그대로 보인다. "
    "컬럼 순서와 표시 항목이 달라 보이는 이유는 대부분 이 때문이다.")).font = BODY
r += 3
r = section(ws, r, "[실측] 행 상태에 따른 입력 가능 여부")
r = hdr(ws, r, ["행 상태", "적합수량", "부적합수량", "부적합사유", "책임사업장", "비고", "회수품입고일자"])
r = row(ws, r, ["미접수·미완료", "가능", "가능", "가능", "가능", "가능", "가능"], BODY, OK_FILL)
r = row(ws, r, ["접수 완료 (STATUS='X')", "잠김", "잠김", "잠김", "잠김", "잠김", "가능"], BODY, NG_FILL)
r += 1
ws.cell(row=r, column=1, value=(
    "접수가 끝난 건은 회수품 입고일자만 고칠 수 있다. 두 프로그램의 잠금 규칙이 같다. "
    "[개발가이드] 품질검사가 마감(ZCRM_QRST_PRT-ZCLOSE='X')된 건도 같은 필드가 잠긴다.")).font = BODY
r += 3
r = put_image(ws, r, "TC01_OLD_2_result.png", "결과화면 예시")
widths(ws, {"A": 24, "B": 26, "C": 14, "D": 16, "E": 14, "F": 12, "G": 18})

# ============================================================ 5. 입력과 저장
ws = wb.create_sheet("5.입력과 저장")
ws["A1"] = "입력과 저장"
ws["A1"].font = TITLE
r = 3
r = section(ws, r, "입력 순서")
r = hdr(ws, r, ["순서", "작업"])
for i, t in enumerate([
    "ALV 에서 대상 행을 찾는다",
    "적합수량 또는 부적합수량을 입력한다",
    "부적합수량을 입력했다면 부적합사유를 선택한다",
    "적합수량을 입력했다면 책임사업장을 입력한다",
    "회수품 입고일자를 입력한다",
    "화면 상단의 접수부서명과 접수자를 입력한다",
    "저장 (Ctrl+S) 을 누른다",
    "'저장 하시겠습니까?' 팝업에서 예를 누른다",
], start=1):
    r = row(ws, r, [i, t])
r += 1
ws.cell(row=r, column=1, value=(
    "[실측 주의] ALV 셀에 값을 넣은 뒤 Enter 를 누르면 화면이 서버 값으로 다시 그려지면서 "
    "입력한 값이 사라진다. 값을 넣었으면 Enter 없이 바로 저장을 누른다.")).font = WARN
r += 3
r = section(ws, r, "부적합사유 코드 (ZCRM_NCF_RSN_T)")
r = hdr(ws, r, ["코드", "내용", "비고"])
rsn = [
    ("Z01", "꼬리표상담번호표시무", ""), ("Z02", "회수기간초과", "실사용 최다"),
    ("Z03", "품명+현물불일치", ""), ("Z04", "파손(검사불가)", ""),
    ("Z05", "꼬리표미부착", ""), ("Z06", "제조일자불일치", ""),
    ("Z07", "교환대상제외", ""), ("Z08", "해피콜부적합미교환", ""),
    ("Z09", "부적합품재청구", ""), ("Z10", "해지대리점", "자동 설정 전용"),
    ("Z11", "테스트모드미진행", ""), ("Z12", "자재번호 없음", ""),
    ("Z13", "중복등록", ""), ("Z14", "조치정보오등록", ""),
    ("Z15", "시스템 에러", ""),
]
for a, b, c2 in rsn:
    r = row(ws, r, [a, b, c2], BODY, OK_FILL if c2 == "실사용 최다" else None)
r += 2
r = section(ws, r, "[실측] 저장 전 업무 검증 6종")
ws.cell(row=r, column=1, value="하나라도 걸리면 저장되지 않고 상태바에 오류(E)가 표시된다.").font = BODY
r += 2
r = hdr(ws, r, ["#", "조건", "메시지"])
val = [
    (1, "적합수량 + 부적합수량 > 요청수량", "적합/부적합 수량의 합은 요청수량을 초과 할 수 없습니다."),
    (2, "부적합수량 > 0 인데 사유가 없음", "부적합 사유가 누락 되었습니다."),
    (3, "부적합수량 = 0 인데 사유가 입력됨", "부적합 수량이 없는데, 부접합 사유가 있습니다."),
    (4, "적합수량 > 0 인데 책임사업장이 없음", "적합 수량이 존재 하면, 책임사업장은 필 수 입니다."),
    (5, "접수부서명 또는 접수자가 없음", "접수부서명 또는 접수자가 누락 되었습니다."),
    (6, "수량을 입력했는데 회수품 입고일자가 없음", "회수품 입고일자를 입력해주세요"),
]
for a, b, c2 in val:
    r = row(ws, r, [a, b, c2], BODY, NG_FILL if a == 6 else None)
r += 1
ws.cell(row=r, column=1, value=(
    "[실측] 6번만 두 프로그램의 동작이 다르다. 기존은 저장 버튼을 누르는 즉시 차단하고 확인 "
    "팝업이 뜨지 않는다. 신규는 확인 팝업을 먼저 띄우고 예를 누른 뒤 차단하며 메시지 끝에 "
    "마침표가 있다.")).font = WARN
r += 3
r = section(ws, r, "[실측] 저장되는 테이블")
ws.cell(row=r, column=1, value=(
    "검증을 통과하면 상태바에 '정상적으로 저장되었습니다.' 가 표시되고 아래 세 테이블이 "
    "순서대로 기록된다.")).font = BODY
r += 2
r = row(ws, r, ["① ZCRM_RECPTNO - 접수번호 채번"], BOLD)
r = hdr(ws, r, ["필드", "저장되는 값"])
r = row(ws, r, ["DATUM", "저장일 (예 2026.09.10)"])
r = row(ws, r, ["SEQNO", "그날의 순번 (1, 2, 3 …)"])
r += 1
ws.cell(row=r, column=1, value="접수번호 = 일자 8자리 + 순번 5자리.  예) 2026.09.10 의 1번 → 2026091000001").font = BOLD
r += 2
r = row(ws, r, ["② ZCRM_RECPT - 접수·적합성 판정 결과"], BOLD)
r = hdr(ws, r, ["필드", "저장되는 값"])
for a, b in [
    ("RECPTNO / RECPTNO_ITEM", "채번된 접수번호 / 품목번호 (10)"),
    ("REQNO / REQNO_ITEM", "회수품 요청번호"),
    ("CSLTNO / CSLTNO_ITEM", "상담번호"),
    ("RECPTDATE", "저장일"),
    ("CF_QTY / NCF_QTY", "입력한 적합수량 / 부적합수량"),
    ("NCF", "적합이면 Y, 부적합이면 N"),
    ("NCF_RSN_V", "선택한 부적합사유 코드"),
    ("RSPC", "입력한 책임사업장"),
    ("STATUS", "X (접수 완료)"),
    ("CONFDAT", "적합 판정일. 적합 저장일 때만 채워진다"),
    ("NETWR / WAERS", "금액 / 통화. 적합 저장일 때만 채워진다"),
    ("REMARK", "입력한 비고"),
    ("REMARK1 / REMARK2", "접수부서명 / 접수자"),
]:
    r = row(ws, r, [a, b])
r += 1
r = row(ws, r, ["③ ZCRM_RECPT_WEAR - 회수품 실제 입고일자"], BOLD)
r = hdr(ws, r, ["필드", "저장되는 값"])
r = row(ws, r, ["RECPTNO / RECPTNO_ITEM", "접수번호 / 품목번호"])
r = row(ws, r, ["WEARDATS", "입력한 회수품 입고일자"])
r += 2
r = section(ws, r, "[실측] 적합 저장과 부적합 저장의 차이")
r = hdr(ws, r, ["필드", "적합수량 저장", "부적합수량 저장"])
for a, b, c2 in [
    ("CF_QTY / NCF_QTY", "1 / 0", "0 / 1"),
    ("NCF", "Y", "N"),
    ("NCF_RSN_V", "공백", "선택한 코드 (예 Z02)"),
    ("RSPC", "필수 입력", "공백 가능"),
    ("CONFDAT", "저장일", "공백"),
    ("NETWR / WAERS", "금액 · KRW", "공백"),
]:
    r = row(ws, r, [a, b, c2])
r += 2
r = section(ws, r, "[개발가이드] 금액 계산")
ws.cell(row=r, column=1, value=(
    "가격 기준 자재는 대체자재(SMATN)가 있으면 그것을, 없으면 MATNR_PRT 를 쓴다. "
    "A911 (VKORG='ASKR', VTWEG='10', KSCHL='PR00', 유효기간) → KONP 에서 단가를 읽어 "
    "NETWR = CF_QTY × KBETR, WAERS = KONWA 로 저장한다.")).font = BODY
r += 3
r = put_image(ws, r, "S10_TC09-01_OLD_4_popup.png", "저장 확인 팝업")
r = put_image(ws, r, "S10_TC09-01_OLD_4b_after_confirm.png", "검증 오류 예시 - 수량 초과")
widths(ws, {"A": 30, "B": 42, "C": 60})

# ============================================================ 6. 상세조회와 종료
ws = wb.create_sheet("6.상세조회와 종료")
ws["A1"] = "상세조회와 화면 종료"
ws["A1"].font = TITLE
r = 3
r = section(ws, r, "회수가능 제조일 상세")
ws.cell(row=r, column=1, value="ALV 에서 한 행만 선택하고 '회수가능 제조일 상세' 또는 Shift+F11").font = BODY
r += 2
r = hdr(ws, r, ["컬럼", "내용"])
for a, b in [
    ("MATNR / MATNRT", "자재 / 자재명"),
    ("PRODH / PRODHT", "제품 계층 / 제품 계층명"),
    ("PRDDATE", "제조일"),
    ("ISTDATE", "설치일"),
    ("PART_PERIOD", "보증기간 (개월)"),
    ("M_FROM", "제조 기준 시작일"),
    ("IST_FROM / IST_TO", "설치 기준 시작일 / 종료일"),
]:
    r = row(ws, r, [a, b])
r += 1
ws.cell(row=r, column=1, value=(
    "[실측] 행을 선택하지 않거나 두 행 이상 선택하면 상태바에 '하나의 라인을 선택 하세요 .' (E) "
    "가 표시되고 팝업이 뜨지 않는다.")).font = BODY
r += 3
r = section(ws, r, "화면 종료")
ws.cell(row=r, column=1, value="뒤로(F3) 를 누른다.").font = BODY
r += 2
r = hdr(ws, r, ["상황", "동작"])
r = row(ws, r, ["아무것도 바꾸지 않음", "팝업 없이 선택화면으로 복귀"], BODY, OK_FILL)
r = row(ws, r, ["값을 바꾸고 저장하지 않음",
                "'저장 하시겠습니까?' 팝업 → 예 저장 / 아니오 버림 / 취소 화면 유지"], BODY, WARN_FILL)
r += 1
ws.cell(row=r, column=1, value="접수부서명이나 접수자만 바꿔도 변경으로 인식된다.").font = BODY
r += 3
r = put_image(ws, r, "TC15_ZCRMR052_onesel.png", "상세조회 팝업")
widths(ws, {"A": 30, "B": 70})

# ============================================================ 7. 자주 겪는 상황
ws = wb.create_sheet("7.자주 겪는 상황")
ws["A1"] = "자주 겪는 상황"
ws["A1"].font = TITLE
r = 3
r = hdr(ws, r, ["증상", "원인과 조치"])
for a, b in [
    ("조회 결과가 0건", "접수일자 기본값이 오늘~오늘이다. 날짜를 넓혀서 조회한다"),
    ("특정 자재가 안 나옴", "[개발가이드] MVKE 에 해당 자재의 ASKR / 10 조합이 없으면 조회되지 않는다"),
    ("조회가 너무 느림", "15일 이내로 끊어서 조회한다. 넓은 범위는 수만 건이 나온다"),
    ("입력한 값이 사라짐", "셀 입력 후 Enter 를 누르면 화면이 다시 그려진다. Enter 없이 저장을 누른다"),
    ("수량을 못 고침", "이미 접수 완료(STATUS='X')된 건이다. 회수품 입고일자만 수정 가능하다"),
    ("필요한 컬럼이 안 보임", "ALV 레이아웃 설정 문제다. 레이아웃을 바꿔 확인한다"),
    ("부적합 사유·책임사업장 버튼이 반응 없음",
     "[실측] 기능이 연결되어 있지 않다. 값은 ALV 셀에서 직접 입력한다"),
]:
    r = row(ws, r, [a, b])
widths(ws, {"A": 36, "B": 88})

# ============================================================ 8. 두 프로그램 차이
ws = wb.create_sheet("8.두 프로그램 차이")
ws["A1"] = "일반 / 구독 전용 프로그램 차이"
ws["A1"].font = TITLE
r = 3
r = hdr(ws, r, ["항목", "ZCRM052 (일반)", "ZCRM052S (구독 전용)"])
for a, b, c2 in [
    ("프로그램", "ZCRMR052", "ZCRMR052S"),
    ("창 제목", "회수품적합성관리", "[구독전용] 회수품 적합성 관리"),
    ("조회 대상", "전체", "ZCRM_CICSVC-SUBS_TYPE = 'S' 인 구독 데이터만"),
    ("ALV 컬럼", "저장된 레이아웃 6개, /CRM01 기본", "레이아웃 없음. 구독구분 컬럼 추가"),
    ("그 외 화면·검증·저장", "-", "동일"),
]:
    r = row(ws, r, [a, b, c2])
r += 1
ws.cell(row=r, column=1, value=(
    "[실측] 2026-09-10 현재 구독 필터가 주석 처리되어 있어 일반 데이터도 조회·저장된다. "
    "비교 테스트를 위한 임시 상태이며 반드시 원복해야 한다.")).font = WARN
r += 3
r = section(ws, r, "확인이 필요한 항목")
r = hdr(ws, r, ["#", "내용"])
r = row(ws, r, [1, "부적합 사유·책임사업장 버튼이 두 프로그램 모두 동작하지 않는다"], BODY, WARN_FILL)
r = row(ws, r, [2, "NCF_RSN_V(부적합사유 코드) 컬럼이 신규 프로그램에 표시되지 않는다"], BODY, WARN_FILL)
r = row(ws, r, [3, "회수품 입고일자 누락 검증의 메시지 문구와 검증 시점이 두 프로그램에서 다르다"],
        BODY, WARN_FILL)
widths(ws, {"A": 26, "B": 44, "C": 54})

# ============================================================ 9. 화면 모음
ws = wb.create_sheet("9.화면 모음")
ws["A1"] = "화면 모음"
ws["A1"].font = TITLE
ws["A2"] = "원본: C:\\sap-test\\results\\shots\\ (PNG 1929x1368)"
ws["A2"].font = NOTE
r = 4
for fn, cap in [
    ("MANUAL_sel_center.png", "선택화면 - 센터조치"),
    ("MANUAL_sel_dealer.png", "선택화면 - 서비스점조치"),
    ("MANUAL_sel_notrecv.png", "선택화면 - 회수품 미접수"),
    ("TC01_OLD_2_result.png", "결과화면 - 일반 프로그램"),
    ("TC01_NEW_2_result.png", "결과화면 - 구독 전용 프로그램"),
    ("S10_TC09-01_OLD_3_entered.png", "ALV 셀 값 입력 상태"),
    ("S10_TC09-01_OLD_4_popup.png", "저장 확인 팝업"),
    ("S10_TC09-01_OLD_4b_after_confirm.png", "검증 오류 - 수량 초과"),
    ("S10_TC07_OLD_4b_after_confirm.png", "저장 완료"),
    ("TC15_ZCRMR052_onesel.png", "상세조회 팝업"),
]:
    r = put_image(ws, r, fn, cap)
widths(ws, {"A": 130})

for sh in wb.worksheets:
    sh.sheet_view.showGridLines = False

wb.save(OUT)
print("saved:", OUT)
print("sheets:", [s.title for s in wb.worksheets])
