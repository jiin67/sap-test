# -*- coding: utf-8 -*-
# build_plain_xlsx.py : DRM 없는 결과서 사본을 openpyxl 로 만든다 (태블릿·모바일 열람용).
#   셀 내용은 dump_xlsx_cells.ps1 이 원본(DRM) 결과서에서 읽은 JSON, 화면 캡처는 페이로드 JSON 의 썸네일 경로를 쓴다.
#   usage: python build_plain_xlsx.py <dump.json> <payload.json> <out.xlsx>
import io, json, math, re, sys
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

dump_p, pay_p, out_p = sys.argv[1], sys.argv[2], sys.argv[3]
dump = json.load(io.open(dump_p, encoding="utf-8-sig"))
pay = json.load(io.open(pay_p, encoding="utf-8"))

FONT = "맑은 고딕"
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
SEC_FONT = Font(name=FONT, bold=True, color="1F4E79", size=11)
TITLE_FONT = Font(name=FONT, bold=True, size=14)
BODY = Font(name=FONT, size=10)
MONO = Font(name="Consolas", size=10)
PASS = PatternFill("solid", fgColor="E2EFDA")
FAIL = PatternFill("solid", fgColor="F8CBAD")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

# 원본 양식의 머리글 행 (시트명 → 행번호들)
HEADER_ROWS = {"개정이력": [4], "공통테스트케이스": [3, 4], "프로그램별테스트케이스": [3, 4], "FAIL내역": [2]}
EVIDENCE = "테스트증적(캡처)"


def is_diagram(text):
    return bool(re.search(r"[│├└┐┌─▼]", text))


def write_sheet(ws, sh):
    rows = sh["rows"]
    heads = set(HEADER_ROWS.get(sh["name"], []))
    free = sh["name"] not in HEADER_ROWS and sh["name"] not in ("표지", "작성가이드")
    prev_section = False
    for r, row in enumerate(rows, start=1):
        filled = [c for c in row if str(c).strip() != ""]
        for c, v in enumerate(row, start=1):
            if v == "":
                continue
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY
            cell.alignment = WRAP
        if not filled:
            prev_section = False
            continue
        first = str(row[0])
        # 표 머리글 / 섹션 제목 / 표 행 꾸미기
        if r in heads:
            for c in range(1, len(row) + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BOX
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            continue
        if free:
            if r == 1 and len(filled) == 1:
                ws.cell(row=r, column=1).font = TITLE_FONT
                continue
            if len(filled) == 1 and re.match(r"^\d+\.\s", first):
                ws.cell(row=r, column=1).font = SEC_FONT
                prev_section = True
                continue
            if is_diagram(first) or first.startswith("  ["):
                ws.cell(row=r, column=1).font = MONO
                ws.cell(row=r, column=1).alignment = Alignment(wrap_text=False, vertical="top")
                continue
            if prev_section and len(filled) > 1:
                for c in range(1, len(filled) + 1):
                    cell = ws.cell(row=r, column=c)
                    cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BOX
                prev_section = False
                continue
            if len(filled) > 1:
                for c in range(1, len(row) + 1):
                    if ws.cell(row=r, column=c).value is not None:
                        ws.cell(row=r, column=c).border = BOX
            prev_section = False
        else:
            maxc = len(row)
            if sh["name"] in HEADER_ROWS and r > max(heads or [0]):
                for c in range(1, maxc + 1):
                    ws.cell(row=r, column=c).border = BOX
    # PASS / FAIL 색
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == "PASS":
                cell.fill = PASS
            elif cell.value == "FAIL":
                cell.fill, cell.font = FAIL, Font(name=FONT, bold=True, size=10)
    for i, w in enumerate(sh["widths"], start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(4, min(float(w), 90))
    for m in sh.get("merges", []):
        try:
            ws.merge_cells(m)
        except Exception:
            pass
    if heads:
        ws.freeze_panes = ws.cell(row=max(heads) + 1, column=1)


def add_picture(ws, path, row, col=2, maxw=700):
    im = PILImage.open(path)
    w, h = im.size
    scale = min(1.0, maxw / w)
    pic = XLImage(path)
    pic.width, pic.height = int(w * scale), int(h * scale)
    ws.add_image(pic, "%s%d" % (get_column_letter(col), row))
    return int(math.ceil(pic.height / 20.0)) + 2


ev_sheets = {EVIDENCE: pay["evidence"]}
for key in ("integ", "integ2", "integ3", "integ4", "integ5"):   # v0.05~: 통합테스트 증적 시트들
    if pay.get(key):
        ev_sheets[pay[key]["ev_sheet"]] = pay[key]["evidence"]

wb = Workbook()
wb.remove(wb.active)
for sh in dump["sheets"]:
    ws = wb.create_sheet(sh["name"][:31])
    if sh["name"] in ev_sheets:
        ws["A1"] = sh["name"]
        ws["A1"].font = TITLE_FONT
        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 110
        r, last = 3, None
        for e in ev_sheets[sh["name"]]:
            if e["tc"] != last:
                ws.cell(row=r, column=1, value=e["tc"]).font = Font(name=FONT, bold=True, size=11)
                last = e["tc"]
            ws.cell(row=r, column=2, value=e["caption"]).font = Font(name=FONT, bold=True, size=10)
            r += 1 + add_picture(ws, e["path"], r + 1)
        continue
    write_sheet(ws, sh)
    # 2026-10-08: 테스트케이스 시트의 표 오른쪽(L열부터)에 그 시트의 증적 캡처를 행별로 넣는다 (원본은 '셀에 배치', 열람용은 떠 있는 그림)
    case_ev = {"프로그램별테스트케이스": pay.get("evidence", [])}
    for key in ("integ", "integ2", "integ3", "integ4", "integ5"):
        if pay.get(key):
            case_ev[pay[key]["sheet"]] = pay[key]["evidence"]
    if sh["name"] in case_ev:
        row_of = {str(rw[0]).strip(): i + 1 for i, rw in enumerate(sh["rows"]) if rw and re.match(r"^[A-Z]-\d{2}$", str(rw[0]).strip())}
        by_tc = {}
        for e in case_ev[sh["name"]]:
            by_tc.setdefault(e["tc"], []).append(e)
        maxn = max([len(v) for v in by_tc.values()] + [1])
        for c in range(12, 12 + maxn):
            ws.column_dimensions[get_column_letter(c)].width = 58
        ws.cell(row=3, column=12, value="증적 캡처 (여러 장이면 오른쪽 열로)").font = Font(name=FONT, bold=True, size=10)
        for tc, evs in by_tc.items():
            rr = row_of.get(tc)
            if not rr:
                continue
            maxh = 0
            for k, e in enumerate(evs):
                im = PILImage.open(e["path"]); w, h = im.size
                scale = 400.0 / w
                pic = XLImage(e["path"]); pic.width, pic.height = int(w * scale), int(h * scale)
                ws.add_image(pic, "%s%d" % (get_column_letter(12 + k), rr))
                maxh = max(maxh, pic.height)
            cur = ws.row_dimensions[rr].height or 15
            ws.row_dimensions[rr].height = max(cur, min(409, maxh * 0.75 + 20))
    if sh["name"] == "FAIL내역":
        # 2026-10-08: 캡처를 각 결함 행의 E열(5열)부터 가로로 나란히 넣는다 (표 아래 방식 → 행 안 방식)
        row_of = {str(rw[0]).strip(): i + 1 for i, rw in enumerate(sh["rows"]) if rw and str(rw[0]).strip()}
        maxn = max([len(f.get("image_paths", [])) for f in pay["fail_rows"]] + [1])
        for c in range(5, 5 + maxn):
            ws.column_dimensions[get_column_letter(c)].width = 58
        for f in pay["fail_rows"]:
            rr = row_of.get(str(f["id"]))
            ips = [ip for ip in f.get("image_paths", []) if ip]
            if not rr or not ips:
                continue
            ws.cell(row=rr, column=5, value="캡처 %d장 (오른쪽)" % len(ips)).font = Font(name=FONT, bold=True, size=10)
            maxh = 0
            for k, ip in enumerate(ips):
                im = PILImage.open(ip); w, h = im.size
                scale = 400.0 / w
                pic = XLImage(ip); pic.width, pic.height = int(w * scale), int(h * scale)
                ws.add_image(pic, "%s%d" % (get_column_letter(5 + k), rr))
                maxh = max(maxh, pic.height)
            ws.row_dimensions[rr].height = min(409, maxh * 0.75 + 20)

wb.save(out_p)
with open(out_p, "rb") as fh:
    sig = fh.read(4)
print("saved", out_p, "sheets", len(wb.sheetnames), "header", sig.hex(), "(504b0304 = 일반 xlsx)")
