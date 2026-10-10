# -*- coding: utf-8 -*-
# make_manual_subs_xlsx.py : step-by-step test manual (일시불구독 개선) with screenshots
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

BASE = r"C:\sap-test\results"
SHOTS = os.path.join(BASE, "shots")
THUMB = os.path.join(SHOTS, "_man_subs")
OUT = os.path.join(BASE, "ZCRM058_일시불구독_테스트절차서_20260912.xlsx")
os.makedirs(THUMB, exist_ok=True)

F, MONO = "Arial", "Consolas"
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
CODE = Font(name=MONO, size=9)
BOLD = Font(name=F, bold=True, size=10)
TITLE = Font(name=F, bold=True, size=15)
SEC = Font(name=F, bold=True, size=12, color="1F3864")
NOTE = Font(name=F, size=9, italic=True, color="666666")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
NG_FILL = PatternFill("solid", fgColor="FCE4D6")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
IN_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
missing = []


def thumb(fn, maxw=900):
    src = os.path.join(SHOTS, fn)
    if not os.path.exists(src):
        missing.append(fn)
        return None
    dst = os.path.join(THUMB, fn)
    im = PILImage.open(src)
    im.thumbnail((maxw, maxw))
    im.save(dst, optimize=True)
    return dst


def hdr(ws, r, cells):
    for j, v in enumerate(cells, start=1):
        c = ws.cell(row=r, column=j, value=v)
        c.fill, c.font = H_FILL, H_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    return r + 1


def row(ws, r, cells, font=BODY, fill=None):
    for j, v in enumerate(cells, start=1):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.border = font, BOX
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if fill:
            c.fill = fill
    return r + 1


def section(ws, r, text):
    ws.cell(row=r, column=1, value=text).font = SEC
    return r + 1


def note(ws, r, text):
    ws.cell(row=r, column=1, value=text).font = NOTE
    return r + 1


def widths(ws, spec):
    for k, v in spec.items():
        ws.column_dimensions[k].width = v


def image(ws, r, fn, caption=None, col="A"):
    p = thumb(fn)
    if not p:
        return r
    if caption:
        c = ws.cell(row=r, column=1, value="■ " + caption)
        c.font = BOLD
        r += 1
    img = XLImage(p)
    ws.add_image(img, "{}{}".format(col, r))
    return r + int(img.height / 19) + 2


wb = Workbook()

# ==================================================== 1. 개요와 준비
ws = wb.active
ws.title = "1.개요와 준비"
widths(ws, {"A": 26, "B": 58, "C": 58, "D": 40})
ws["A1"] = "AS 무상지원금 – 일시불구독 개선 테스트 절차서"
ws["A1"].font = TITLE
ws["A2"] = "직접 따라 하기 위한 안내. 2026-09-11 KDQ 210 에서 실제로 수행한 절차와 값을 그대로 적었다."
ws["A2"].font = NOTE
r = 4
r = section(ws, r, "1) 대상")
r = hdr(ws, r, ["항목", "내용", "비고"])
for a, b, c in [
    ("메인 프로그램", "ZCRMR058_1", "T-code 058A(일반) / 058S(구독) / 059A / 059S(구독 조회전용)"),
    ("상세 프로그램", "ZCRMR060_1", "T-code 060A(일반) / 060S(구독)"),
    ("신규 RFC", "ZCRMF03_RFC_ZCRM_CHRG_CNFRM", "레거시(CIC)가 상담건별 확정일을 보내는 창구"),
    ("신규 테이블", "ZCRM_CHRG_CNFRM", "상담번호 / 월 / 대리점확정일 + 생성·변경 이력"),
    ("구분자 필드", "ZCRM_CICSVC_NEW-SUBS_TYPE", "'S' = 구독. SE16N 설명은 '구독여부 지시자'"),
    ("시스템", "KDQ 클라이언트 210", "KDP(운영)에서는 절대 수행하지 않는다"),
    ("기간월", "2026.07", "QA 에 데이터가 있는 마지막 달. 2026.08 은 0건이다"),
]:
    r = row(ws, r, [a, b, c])
r += 1

r = section(ws, r, "2) 이 테스트의 핵심 로직 (전건확정)")
r = row(ws, r, ["구독 건의 대리점확정일(CFDATE1)은 화면에서 입력하지 않는다. 대리점의 대상 상담건이 "
                "ZCRM_CHRG_CNFRM 에 '전부' 확정돼 있을 때만 확정일(최대값)이 표시되고, 한 건이라도 빠지면 공백이다. "
                "그래서 테스트는 '확정 건수를 0 → 2 → 3 → 2 로 바꿔가며 화면의 CFDATE1 을 본다' 가 전부다.", None, None])
r += 2

r = section(ws, r, "3) 미리 알아둘 점 (여기서 직접 겪은 것)")
r = hdr(ws, r, ["주의", "내용"])
for a, b in [
    ("필수 입력", "058/059 선택화면은 기간월과 센터를 모두 넣어야 실행된다. 하나라도 비면 "
                "'모든 필수 입력 필드에 값을 입력하십시오.' 가 뜬다"),
    ("SE16N 편집", "엔트리 유지보수(편집) 체크가 켜진 채 남아 있는 경우가 잦다. 시작할 때 확인하고, "
                 "테스트가 끝나면 반드시 끈다"),
    ("변경이력", "SE16N 편집으로 바꾼 값은 AENAM/AEDAT(변경자·변경일)를 남기지 않는다. "
              "그래서 바꾼 상담번호 목록을 따로 적어둬야 원복할 수 있다"),
    ("전표 버튼", "'일괄회계전표생성'은 '계속 진행 하시겠습니까?' 팝업에서 [예]를 누르면 실제 회계전표가 생긴다. "
               "QA 라도 개발자와 합의 없이 누르지 않는다"),
    ("원복", "테스트가 끝나면 ZCRM_CHRG_CNFRM 행 삭제 + SUBS_TYPE 공백 복원까지 해야 원래 상태가 된다"),
]:
    r = row(ws, r, [a, b], BODY, WARN_FILL)
r += 1

r = section(ws, r, "4) 준비물")
r = hdr(ws, r, ["항목", "내용"])
for a, b in [
    ("권한", "SE16N(조회+편집), SE37, ZCRM058A/058S/059S/060A/060S"),
    ("로그온", "SAP Logon 으로 KDQ 210 에 직접 로그온한다"),
    ("소요 시간", "수동으로 전 과정 약 1.5~2시간 (데이터 세팅 20분, 그룹1 20분, 그룹2 30분, 그룹3 30분, 원복 15분)"),
    ("기록", "각 단계에서 화면을 캡처해 둔다. 판정은 화면 제목이 아니라 금액·확정일 값으로 한다"),
]:
    r = row(ws, r, [a, b])

# ==================================================== 2. 테스트 데이터 만들기
ws = wb.create_sheet("2.테스트데이터 만들기")
widths(ws, {"A": 10, "B": 30, "C": 62, "D": 52})
ws["A1"] = "테스트 데이터 만들기"
ws["A1"].font = TITLE
ws["A2"] = "구독(SUBS_TYPE='S') 데이터가 QA 에 없으므로, 일반 상담건 몇 건을 잠시 구독으로 바꿔 쓴다."
ws["A2"].font = NOTE
r = 4

r = section(ws, r, "STEP 1. 후보 상담건 찾기 (SE16N, 조회만)")
r = hdr(ws, r, ["단계", "화면", "입력/조작", "기대 결과"])
for a, b, c, d in [
    ("1-1", "/nSE16N", "테이블 = ZCRM_CICSVC_NEW, 최대 적중 수 = 1500, 엔트리 유지보수 체크 해제", "조회 조건 목록이 뜬다"),
    ("1-2", "SE16N 조건", "조치일자 = 2026.07.01 ~ 2026.07.03\n유무상구분자 = N\n상담번호의 seq(행)번호 = 1",
     "조건 3개가 입력된다"),
    ("1-3", "실행(F8)", "-", "1,500건(최대치) 조회"),
    ("1-4", "결과 ALV", "아래 조건으로 후보를 고른다:\n"
                      "· 클레임유형(CLMART) = A/B/C/D/D1/D2\n"
                      "· 제외사유(DRRSN) 공백\n"
                      "· 기사명(TECHNAME) 있음\n"
                      "· 구독여부 지시자(SUBS_TYPE) 공백\n"
                      "· 서비스점(KUNNR_SVC)이 CT 로 시작하지 않음(센터 제외)",
     "2026-09-11 기준 574건이 조건을 만족했다"),
    ("1-5", "집계", "서비스점(KUNNR_SVC)별로 묶어 3건 이상인 대리점 = X 후보, 1건인 대리점 = Y 후보",
     "X 후보 117786(17건), Y 후보 100011(1건)"),
]:
    r = row(ws, r, [a, b, c, d])
r += 1
r = note(ws, r, "※ 대리점은 KUNNR 가 아니라 KUNNR_SVC(서비스점)다. KUNNR 은 고객코드다.")
r += 1

r = section(ws, r, "STEP 2. 대리점이 쓸 수 있는지 확인 (SE16N, 조회만)")
r = hdr(ws, r, ["단계", "테이블", "조건", "확인 내용"])
for a, b, c, d in [
    ("2-1", "KNVV", "고객 = 117786 / 100011",
     "VKORG=ASKR 행이 있어야 하고, 그 행의 영업소(VKBUR)를 적어둔다. 058 실행 시 '센터'에 넣을 값이다 "
     "(117786 → CT11, 100011 → CT02)"),
    ("2-2", "ZCRM_ASSPTFEEH1", "정산기간 = 2026.07",
     "고르려는 대리점 행이 '없어야' 한다. 있으면 이미 확정/전표가 걸린 달이다. "
     "2026.06 까지는 전표(BELNR)가 다 차 있고 2026.07 은 비어 있었다"),
]:
    r = row(ws, r, [a, b, c, d])
r += 1

r = section(ws, r, "STEP 3. 실제로 사용한 데이터")
r = hdr(ws, r, ["역할", "대리점(KUNNR_SVC)", "상담번호(CSLTNO)", "비고"])
for a, b, c, d in [
    ("X-1", "117786 제주진영서비스점 / 센터 CT11", "2026070100278", "조치일자 2026.07.01, 클레임유형 D1, 기사 0000130784"),
    ("X-2", "117786", "2026070100450", "조치일자 2026.07.02, D1, 기사 0000132094"),
    ("X-3", "117786", "2026070100862", "조치일자 2026.07.01, D1, 기사 0000134384"),
    ("X-예비", "117786", "2026070101452", "쓰지 않았다. 필요하면 예비로"),
    ("Y-1", "100011 (주)승보상사 / 센터 CT02", "2026070102721", "조치일자 2026.07.02, 클레임유형 A, 기사 0000130144"),
]:
    r = row(ws, r, [a, b, c, d], BODY, IN_FILL)
r += 1
r = note(ws, r, "※ 모두 CSLTNO_ITEM=1, 유무상=N, 원래 SUBS_TYPE 은 공백이다. 원복할 때 되돌릴 값이 '공백'이라는 뜻이다.")
r += 1

r = section(ws, r, "STEP 4. SUBS_TYPE 을 'S' 로 바꾸기 (SE16N 편집, 4건 반복)")
r = row(ws, r, ["주의", "SE16N 편집은 테이블을 직접 고치는 기능이다. 아래 상담번호 4건 외에는 절대 건드리지 않는다. "
                     "한 건씩, 조회 결과가 1행인지 확인한 뒤에 고친다.", None, None], BODY, WARN_FILL)
r += 1
r = hdr(ws, r, ["단계", "화면", "입력/조작", "기대 결과"])
for a, b, c, d in [
    ("4-1", "/nSE16N", "테이블 = ZCRM_CICSVC_NEW, 최대 적중 수 = 10", "-"),
    ("4-2", "SE16N", "'엔트리 유지보수' 체크 ON", "편집 모드로 조회된다"),
    ("4-3", "SE16N 조건", "상담번호 = 2026070100278, 상담번호의 seq(행)번호 = 1", "-"),
    ("4-4", "실행(F8)", "-", "1행만 조회. 구독여부 지시자(SUBS_TYPE) 칸이 비어 있고 입력 가능하다"),
    ("4-5", "결과 ALV", "구독여부 지시자 칸에 S 입력 → 저장(Ctrl+S)", "'변경 완료' 팝업: 삭제 0 / 삽입 0 / 변경 1"),
    ("4-6", "반복", "2026070100450, 2026070100862, 2026070102721 에 대해 4-3~4-5 반복", "4건 모두 변경 1"),
    ("4-7", "마무리", "'엔트리 유지보수' 체크 해제", "편집 모드 해제"),
]:
    r = row(ws, r, [a, b, c, d])
r += 1
r = image(ws, r, "S13_20260911_2301_2026070100278_before.png", "4-4 편집 모드 조회 결과 (1행, SUBS_TYPE 공백)")
r = image(ws, r, "S13_20260911_2301_popup_save_2026070100278.png", "4-5 저장 후 '변경 완료' 팝업 (변경된 행 1)")
r += 1

r = section(ws, r, "STEP 5. 세팅 확인")
r = hdr(ws, r, ["단계", "화면", "입력/조작", "기대 결과"])
r = row(ws, r, ["5-1", "/nSE16N", "테이블 = ZCRM_CICSVC_NEW, 조건 '구독여부 지시자' = S, 편집 체크 해제, 실행",
                "정확히 4건만 나와야 한다 (위 표의 상담번호)"], BODY, OK_FILL)
r += 1
r = note(ws, r, "※ 4건보다 많이 나오면 다른 사람이 만든 구독 데이터가 있다는 뜻이다. 그 건은 건드리지 말고 대상에서 제외한다.")

# ==================================================== 3. 그룹1 화면 분리
ws = wb.create_sheet("3.그룹1 화면분리")
widths(ws, {"A": 10, "B": 26, "C": 52, "D": 60})
ws["A1"] = "그룹 1. 화면 분리 (T-code 별로 구독/일반이 갈리는지)"
ws["A1"].font = TITLE
ws["A2"] = "선택화면 입력은 모두 동일: 기간월 2026.07 / 센터 CT11 / 대리점 117786 (Y 는 센터 CT02, 대리점 100011)"
ws["A2"].font = NOTE
r = 4
r = hdr(ws, r, ["ID", "실행", "입력값", "기대 결과 / 2026-09-11 실제 값"])
for a, b, c, d, f in [
    ("1-1", "/nZCRM058S", "기간월 2026.07, 센터 CT11, 대리점 117786 → 실행(F8)",
     "구독 3건만 집계. 1행, 무상지원금액 20,000 (심야 10,000 + 주간 10,000). 세팅 전에는 0행이었다", OK_FILL),
    ("1-2", "/nZCRM058A", "동일",
     "구독 건이 빠진 금액. 170,000 (세팅 전 190,000 → −20,000). Y 는 60,000 (세팅 전 65,000 → −5,000)", OK_FILL),
    ("1-3", "/nZCRM059S", "동일",
     "구독 1행 20,000 이 보이고, 툴바에 지급대상데이터조회·상세리스트조회·회계전표출력만 있다. "
     "무상지원금확정/취소, 대리점 확정/확정취소, 일괄회계전표생성/삭제 없음", OK_FILL),
    ("1-4", "/nZCRM060S", "기간월 2026.07(SO_SPMON), 센터 CT11, 대리점 117786 → 실행",
     "구독 3건의 기사 3명(0000130784 김덕경 / 0000132094 김태현 / 0000134384 고대의), 합계 20,000. "
     "060A 는 기사 8명 합계 170,000", OK_FILL),
    ("1-5", "058S 결과화면", "행을 선택하고 툴바 '상세리스트조회'(Ctrl+Shift+F7)",
     "ZCRMR060_1 상세로 이동. 060S 와 같은 7행, 합계 20,000", OK_FILL),
]:
    r = row(ws, r, [a, b, c, d], BODY, f)
r += 1
r = note(ws, r, "※ 060 결과화면은 행이 있어도 제목의 건수가 (0) 으로 보인다. 060A 도 같아서 이번 개선과는 무관한 기존 동작으로 보인다.")
r += 1
r = image(ws, r, "SUBS_pre_058S_X.png", "[비교] 세팅 전 ZCRM058S 117786 - 0행")
r = image(ws, r, "SUBS_post_058S_X.png", "1-1 세팅 후 ZCRM058S 117786 - 1행 20,000원")
r = image(ws, r, "SUBS_pre_058A_X.png", "[비교] 세팅 전 ZCRM058A 117786 - 190,000원")
r = image(ws, r, "SUBS_post_058A_X.png", "1-2 세팅 후 ZCRM058A 117786 - 170,000원")
r = image(ws, r, "SUBS_post_059S_X.png", "1-3 ZCRM059S - 처리 버튼이 없다")
r = image(ws, r, "SUBS_post_060S_X.png", "1-4 ZCRM060S - 구독 3건의 기사 3명, 합계 20,000")
r = image(ws, r, "SUBS_post_058S_X_drill_drill.png", "1-5 058S → 상세리스트조회 드릴다운 결과")

# ==================================================== 4. 그룹2 RFC
ws = wb.create_sheet("4.그룹2 RFC")
widths(ws, {"A": 10, "B": 30, "C": 50, "D": 62})
ws["A1"] = "그룹 2. RFC 단독 테스트 (SE37)"
ws["A1"].font = TITLE
ws["A2"] = "ZCRMF03_RFC_ZCRM_CHRG_CNFRM 을 레거시 없이 SE37 에서 직접 호출한다."
ws["A2"].font = NOTE
r = 4
r = section(ws, r, "1) SE37 조작 방법")
r = hdr(ws, r, ["단계", "화면", "입력/조작", "기대 결과"])
for a, b, c, d in [
    ("1", "/nSE37", "Function 모듈 = ZCRMF03_RFC_ZCRM_CHRG_CNFRM → 실행(F8)", "'Function 모듈테스트: 초기화면'"),
    ("2", "테스트 화면", "Import 매개변수 IV_ACTION 의 '값' 칸에 S(등록/변경) 또는 C(취소) 입력", "-"),
    ("3", "테스트 화면", "IT_DATA 줄의 값 칸에 커서를 두고 F2(선택)", "'Structure Editor: 변경 IT_DATA' 로 들어간다"),
    ("4", "Structure Editor", "CSLTNO 칸에 상담번호, CFDATE1 칸에 확정일(YYYYMMDD) 입력. "
                            "줄을 더 넣으려면 툴바 '신규라인'(Ctrl+F8)", "입력한 줄 수만큼 표시"),
    ("5", "Structure Editor", "Enter 후 F3(뒤로)", "테스트 화면의 IT_DATA 옆에 'n 엔트리' 로 표시된다"),
    ("6", "테스트 화면", "실행(F8)", "'Function 모듈 테스트: 결과화면'"),
    ("7", "결과화면", "Export 매개변수 EV_RESULT 값 확인. ET_RETURN 줄에 커서를 두고 F2",
     "건별 CSLTNO / T(유형) / MESSAGE 를 볼 수 있다"),
]:
    r = row(ws, r, [a, b, c, d])
r += 1
r = image(ws, r, "SE37_test_screen.png", "SE37 테스트 화면 (IV_ACTION / IT_DATA)")
r = image(ws, r, "S14_TC2-1_1_it_data.png", "IT_DATA 입력 (CSLTNO / CFDATE1)")
r = image(ws, r, "S14_TC2-1_4_et_return.png", "ET_RETURN 확인 화면")
r += 1

r = section(ws, r, "2) 케이스별 입력값과 결과 (2026-09-11 실측)")
r = hdr(ws, r, ["ID", "IV_ACTION / IT_DATA", "기대 결과", "실제 결과"])
for a, b, c, d, f in [
    ("2-1", "S / 2026070100862:20260807, 2026070102721:20260804, 2026070100278:20260805",
     "EV_RESULT='S', 3건 저장", "EV_RESULT=E. 3건 모두 '존재하지 않는 상담번호'. 저장 0건 → FAIL", NG_FILL),
    ("2-2", "C / 2026070100278, 2026070100450", "EV_RESULT='S', 2건 삭제",
     "EV_RESULT=E, '존재하지 않는 상담번호'. 삭제 0건 → FAIL", NG_FILL),
    ("2-3", "공백 / 9999999999999:20260810", "'E', 처리구분 미입력",
     "SE37 이 필수 매개변수 공백 실행을 막는다('매개변수값 입력' 팝업) → SE37 로는 확인 불가", WARN_FILL),
    ("2-4", "X / 9999999999999:20260810", "'E', 유효하지 않은 구분",
     "EV_RESULT=E, '유효하지 않은 처리구분입니다 (S 또는 C)' → PASS", OK_FILL),
    ("2-5", "S / 2026070100278:20260805, 9999999999999:20260805", "없는 건만 'E', 유효건은 저장",
     "유효건까지 '존재하지 않는 상담번호' → FAIL", NG_FILL),
    ("2-6", "S / 2026070100450:(확정일 공백)", "'E' 확정일자 누락",
     "'E' 이지만 사유가 '존재하지 않는 상담번호' → FAIL", NG_FILL),
    ("2-7", "S / 2026070100450:20261231", "'E' 미래일자 불가",
     "'E' 이지만 사유가 '존재하지 않는 상담번호' → FAIL", NG_FILL),
    ("2-8", "S / (IT_DATA 없음)", "'E', 전송 데이터 없음", "EV_RESULT=E, '전송 데이터가 없습니다' → PASS", OK_FILL),
    ("2-9", "S / 2026070100862:20260807 두 줄", "중복 제거되어 1건 저장",
     "두 줄 모두 '존재하지 않는 상담번호'. 저장 0건 → FAIL", NG_FILL),
]:
    r = row(ws, r, [a, b, c, d], BODY, f)
r += 1

r = section(ws, r, "3) 결함 확인 방법 (중요)")
r = row(ws, r, ["ET_RETURN 에 돌아온 CSLTNO 를 입력값과 비교하면 3자리가 밀려 있다. "
                "CSLTNO(13자리)와 CFDATE1(8자리)을 이어 붙인 문자열의 4~16번째 문자가 CSLTNO 로 쓰인다. "
                "직접 확인하려면 2-5 를 넣고 ET_RETURN 의 CSLTNO 를 그대로 읽어보면 된다.", None, None, None])
r += 1
r = hdr(ws, r, ["보낸 CSLTNO", "보낸 CFDATE1", "RFC 가 받은 CSLTNO", "메시지"])
for a, b, c, d in [
    ("2026070100278", "20260805", "6070100278202", "존재하지 않는 상담번호"),
    ("2026070100450", "(공백=00000000)", "6070100450000", "존재하지 않는 상담번호"),
    ("2026070100450", "20261231", "6070100450202", "존재하지 않는 상담번호"),
    ("9999999999999", "20260805", "9999999999202", "존재하지 않는 상담번호"),
]:
    r = row(ws, r, [a, b, c, d], CODE, NG_FILL)
r += 1
r = note(ws, r, "※ SE37 입력 화면의 IT_DATA 행에는 CSLTNO / CFDATE1 두 필드뿐이고 MANDT(클라이언트)가 없다. "
                "즉 입력은 정상이고, 함수 안에서 MOVE-CORRESPONDING 대신 단순 MOVE 로 "
                "ZCRM_CHRG_CNFRM 형(MANDT 3자리가 맨 앞) 작업영역에 옮길 때 생기는 증상으로 보인다.")

# ==================================================== 5. 그룹3 전건확정
ws = wb.create_sheet("5.그룹3 전건확정")
widths(ws, {"A": 10, "B": 34, "C": 50, "D": 58})
ws["A1"] = "그룹 3. 전건확정 표시 로직 (핵심)"
ws["A1"].font = TITLE
ws["A2"] = "원래는 RFC 가 확정 데이터를 넣어야 하지만 위 결함으로 불가능하다. 화면 로직만 보려면 ZCRM_CHRG_CNFRM 에 직접 넣는다."
ws["A2"].font = NOTE
r = 4
r = row(ws, r, ["주의", "이 방법은 테스트 목적의 우회다. 넣는 행의 생성자(ERNAM)에 ZTEST0911 처럼 표시를 남겨두면 "
                     "나중에 지울 때 헷갈리지 않는다. 테스트가 끝나면 넣은 행을 모두 지운다.", None, None], BODY, WARN_FILL)
r += 2
r = section(ws, r, "1) 확정 행 넣는 방법 (SE16N 편집)")
r = hdr(ws, r, ["단계", "화면", "입력/조작", "기대 결과"])
for a, b, c, d in [
    ("1", "/nSE16N", "테이블 = ZCRM_CHRG_CNFRM, 최대 적중 수 = 50, '엔트리 유지보수' 체크 ON", "-"),
    ("2", "실행(F8)", "-", "0건이어도 편집 가능한 빈 ALV 가 열린다"),
    ("3", "결과 ALV", "툴바의 '행 추가'(Append row) 클릭", "빈 줄이 하나 생긴다"),
    ("4", "결과 ALV", "상담번호 / 월(2026.07) / 대리점확정일 / 생성자(ZTEST0911) 입력", "-"),
    ("5", "저장", "Ctrl+S", "'변경 완료' 팝업의 삽입 건수가 넣은 줄 수와 같아야 한다"),
    ("6", "마무리", "'엔트리 유지보수' 체크 해제", "-"),
]:
    r = row(ws, r, [a, b, c, d])
r += 1
r = image(ws, r, "SE16N_edit_empty.png", "2단계: 0건이어도 편집 그리드가 열린다")
r = image(ws, r, "S15_stepA_after_save.png", "5단계: 확정 2건 입력 후 (생성자 ZTEST0911)")
r += 1

r = section(ws, r, "2) 단계별로 화면 확인하기 (058S, 기간월 2026.07)")
r = hdr(ws, r, ["ID", "ZCRM_CHRG_CNFRM 상태", "확인할 화면", "기대 / 실제 결과"])
for a, b, c, d, f in [
    ("3-1", "0건", "058S 센터 CT11 / 대리점 117786", "CFDATE1 공백 → PASS", OK_FILL),
    ("3-2", "X-1 2026.08.05, X-2 2026.08.06 (2/3)", "058S 117786",
     "공백 유지 (한 건이라도 빠지면 표시 안 함) → PASS", OK_FILL),
    ("3-3", "+ X-3 2026.08.07 (3/3)", "058S 117786", "CFDATE1 = 2026.08.07 (확정일 최대값) → PASS", OK_FILL),
    ("3-4", "X-3 행 삭제 (2/3)", "058S 117786", "다시 공백 → PASS", OK_FILL),
    ("3-5", "Y-1 2026.08.04, 월 2026.07 (1/1)", "058S 센터 CT02 / 대리점 100011",
     "CFDATE1 = 2026.08.04 (대상 1건이면 즉시 표시) → PASS", OK_FILL),
    ("5-2", "Y-1 의 월을 2026.06 으로 넣었을 때", "058S 100011 (기간월 2026.07)",
     "공백. 월을 2026.07 로 고치면 2026.08.04 가 표시된다 → PASS", OK_FILL),
    ("회귀", "X 3/3 확정 상태", "058A 117786",
     "058A 의 CFDATE1 은 공백 그대로 (구독 확정일이 일반 화면에 새지 않는다) → PASS", OK_FILL),
]:
    r = row(ws, r, [a, b, c, d], BODY, f)
r += 1
r = image(ws, r, "SUBS_t32_058S_X.png", "3-2 부분확정(2/3) - CFDATE1 공백")
r = image(ws, r, "SUBS_t33_058S_X.png", "3-3 전건확정(3/3) - CFDATE1 2026.08.07")
r = image(ws, r, "SUBS_t34_058S_X.png", "3-4 1건 취소 후 - 다시 공백")
r = image(ws, r, "SUBS_t35_058S_Y.png", "3-5 단일건 대리점 - 2026.08.04")
r = image(ws, r, "SUBS_t52n_058S_Y.png", "5-2 월이 2026.06 인 확정 행은 2026.07 조회에 잡히지 않는다")
r = image(ws, r, "SUBS_t33_058A_X.png", "회귀: 같은 시점의 058A - CFDATE1 공백")

# ==================================================== 6. 그룹4·5
ws = wb.create_sheet("6.그룹4·5 버튼과 원복")
widths(ws, {"A": 10, "B": 30, "C": 52, "D": 58})
ws["A1"] = "그룹 4·5. 버튼 제어 / 전표 / 원복"
ws["A1"].font = TITLE
r = 3
r = section(ws, r, "1) 버튼 제어")
r = hdr(ws, r, ["ID", "화면", "확인 방법", "결과"])
for a, b, c, d, f in [
    ("4-1", "ZCRM058S 결과화면", "툴바에서 '대리점 확정'(F2)·'대리점 확정취소'(F5) 를 찾는다",
     "없다 → PASS (구독은 수동 확정을 막는 것이 목적)", OK_FILL),
    ("4-2", "ZCRM058A 결과화면", "같은 버튼을 찾는다", "있고 활성 → PASS (일반은 기존대로 수동 확정)", OK_FILL),
]:
    r = row(ws, r, [a, b, c, d], BODY, f)
r += 1
r = image(ws, r, "SUBS_pre_058A_X.png", "4-2 058A 툴바 - 대리점 확정 / 대리점 확정취소 있음")
r = image(ws, r, "SUBS_post_058S_X.png", "4-1 058S 툴바 - 두 버튼이 없다")
r += 1

r = section(ws, r, "2) 전표 (4-3 / 4-4) — 이번에 실행하지 않았다")
r = hdr(ws, r, ["ID", "내용", "하지 않은 이유", "하려면"])
for a, b, c, d in [
    ("4-3", "전건확정 후 일괄회계전표생성 → 전표 생성",
     "KDQ 에 실제 회계전표가 생긴다. 또 ZCRM_ASSPTFEEH1 의 키가 기간월+대리점이라 같은 대리점의 일반건과 행을 공유한다",
     "개발자와 함께, 전표 삭제(일괄회계전표삭제)까지 계획하고 진행"),
    ("4-4", "부분확정 상태에서 전표생성 → '확정일자 없음' 차단 확인",
     "버튼을 누르면 '계속 진행 하시겠습니까?' 팝업이 확정일자 검사보다 먼저 뜬다. [예]를 눌렀을 때 차단이 동작하지 않으면 전표가 생긴다",
     "3-2 상태(2/3)에서 058S 117786 조회 → 일괄회계전표생성 → [예]. 차단 메시지가 나오면 PASS, "
     "전표가 생기면 FAIL 이므로 즉시 전표 삭제 필요"),
]:
    r = row(ws, r, [a, b, c, d], BODY, WARN_FILL)
r += 1
r = image(ws, r, "SUBS_t44_058S_X_fipos_popup1.png", "4-4 일괄회계전표생성 클릭 시 먼저 뜨는 확인 팝업")
r += 1

r = section(ws, r, "3) 원복 (5-3) — 반드시 수행")
r = hdr(ws, r, ["단계", "대상", "방법", "확인"])
for a, b, c, d in [
    ("1", "ZCRM_CHRG_CNFRM", "SE16N 편집으로 넣은 행(생성자 ZTEST0911)을 선택해 '행 삭제' 후 저장", "조회하면 '값이 없습니다'"),
    ("2", "ZCRM_CICSVC_NEW-SUBS_TYPE", "4건을 다시 공백으로 (S → 공백, 같은 방법)",
     "조건 '구독여부 지시자'=S 로 조회하면 0건"),
    ("3", "화면 확인", "058A 117786 / 100011, 058S 양쪽, 060S 117786",
     "058A 190,000 / 65,000 (세팅 전과 동일), 058S·060S 0행"),
]:
    r = row(ws, r, [a, b, c, d], BODY, OK_FILL)
r += 1
r = image(ws, r, "SUBS_r53_058A_X.png", "원복 후 058A 117786 - 190,000 (세팅 전과 같다)")
r = image(ws, r, "SUBS_r53_058S_X.png", "원복 후 058S 117786 - 0행")

# ==================================================== 7. 스크립트로 재현
ws = wb.create_sheet("7.스크립트로 재현")
widths(ws, {"A": 46, "B": 74, "C": 46})
ws["A1"] = "스크립트로 같은 절차 재현하기 (선택)"
ws["A1"].font = TITLE
ws["A2"] = "수동 대신 쓸 수 있다. SAP 에 직접 로그온한 상태에서 PowerShell 로 실행하며, 스크립트가 별도 세션을 열어 작업한다."
ws["A2"].font = NOTE
r = 4
r = hdr(ws, r, ["스크립트", "실행 예", "하는 일"])
for a, b, c in [
    ("scripts\\S13_se16n_subs_type.vbs",
     "cscript //nologo scripts\\S13_se16n_subs_type.vbs dry S 2026070100278\n"
     "cscript //nologo scripts\\S13_se16n_subs_type.vbs apply S 2026070100278 2026070100450 2026070100862 2026070102721\n"
     "cscript //nologo scripts\\S13_se16n_subs_type.vbs apply _ (같은 4건)",
     "SUBS_TYPE 세팅/원복. dry 는 조회만 하고 바꾸지 않는다. 1행·기존값 확인 후에만 저장한다"),
    ("scripts\\S15_se16n_chrg_cnfrm.vbs",
     "cscript //nologo scripts\\S15_se16n_chrg_cnfrm.vbs stepA ins:2026070100278:2026.07:2026.08.05 "
     "ins:2026070100450:2026.07:2026.08.06\n"
     "(다른 예) del:2026070100862 / upd:2026070102721:2026.07:2026.08.04",
     "ZCRM_CHRG_CNFRM 확정 행 추가/수정/삭제. 지정한 4개 상담번호 외에는 거부한다"),
    ("scripts\\S14_se37_rfc.vbs",
     "cscript //nologo scripts\\S14_se37_rfc.vbs TC2-5 S 2026070100278:20260805 9999999999999:20260805",
     "SE37 로 RFC 호출. 실행 전에 IV_ACTION 과 엔트리 수를 검사하고, 결과화면과 ET_RETURN 을 로그로 남긴다"),
    ("scripts\\inspect_058_result.vbs",
     "cscript //nologo scripts\\inspect_058_result.vbs ZCRM058S 117786 2026.07 mytag CT11\n"
     "(끝에 DRILL 을 붙이면 상세리스트조회까지, FIPOS 를 붙이면 전표 버튼까지)",
     "058/059/060 을 실행해 ALV 값·툴바·화면 캡처를 남긴다. 읽기 전용"),
    ("scripts\\S12_se16n_cols.vbs",
     "cscript //nologo scripts\\S12_se16n_cols.vbs ZCRM_CICSVC_NEW 50 tag \"CSLTNO,SUBS_TYPE\" \"구독여부 지시자=S\"",
     "SE16N 조회 결과를 TSV 로 덤프 (조회 전용)"),
]:
    r = row(ws, r, [a, b, c], CODE)
r += 1
r = hdr(ws, r, ["항목", "경로", "비고"])
for a, b, c in [
    ("실행 로그", "C:\\sap-test\\logs\\", "UTF-16. PowerShell 에서 Get-Content -Encoding Unicode 로 읽는다"),
    ("화면 캡처", "C:\\sap-test\\results\\shots\\", "PNG"),
    ("결과 요약", "C:\\sap-test\\results\\ZCRM058_일시불구독_20260911.md", "이번 테스트의 상세 기록"),
    ("결과서", "C:\\sap-test\\results\\SAP_KR_단위테스트케이스및결과서_ZCRMR058_1_일시불구독_v0.01_20260911.xlsx", "회사 양식"),
]:
    r = row(ws, r, [a, b, c])
r += 1
r = note(ws, r, "※ 스크립트는 KDQ 가 아니면 즉시 종료한다. 장시간 실행 중인 스크립트를 강제 종료하면 "
                "SAP 스크립팅 엔진 연결이 끊겨 SAP 를 다시 띄워야 하므로, 끝날 때까지 기다린다.")

# ==================================================== 8. 캡처 모음
ws = wb.create_sheet("8.캡처 모음")
widths(ws, {"A": 20, "B": 120})
ws["A1"] = "화면 캡처 모음"
ws["A1"].font = TITLE
r = 3
for tag, fn, cap in [
    ("데이터 세팅", "S13_20260911_2301_2026070100278_before.png", "SE16N 편집 조회 (SUBS_TYPE 공백)"),
    ("데이터 세팅", "S13_20260911_2334_2026070100278_after.png", "저장 후"),
    ("1-1", "SUBS_post_058S_X.png", "ZCRM058S 117786 - 20,000"),
    ("1-1", "SUBS_post_058S_Y.png", "ZCRM058S 100011 - 5,000"),
    ("1-2", "SUBS_post_058A_X.png", "ZCRM058A 117786 - 170,000"),
    ("1-2", "SUBS_post_058A_Y.png", "ZCRM058A 100011 - 60,000"),
    ("1-3", "SUBS_post_059S_X.png", "ZCRM059S 117786"),
    ("1-4", "SUBS_post_060A_X.png", "ZCRM060A 117786 - 170,000"),
    ("2-1", "S14_TC2-1_3_result.png", "RFC 2-1 결과화면 (EV_RESULT=E)"),
    ("2-4", "S14_TC2-4_4_et_return.png", "RFC 2-4 ET_RETURN - 유효하지 않은 처리구분"),
    ("2-8", "S14_TC2-8_4_et_return.png", "RFC 2-8 ET_RETURN - 전송 데이터가 없습니다"),
    ("2-3", "S14_TC2-3_popup.png", "RFC 2-3 - 매개변수값 입력 팝업"),
    ("3-2", "S15_stepA_after_save.png", "확정 2건 입력 상태"),
    ("3-3", "S15_stepB_after_save.png", "확정 3건 + Y(2026.06) 입력 상태"),
    ("3-4", "S15_stepD_after_save.png", "X-3 삭제 후"),
    ("5-3", "S15_cleanup_after_save.png", "확정 행 전체 삭제 후"),
    ("5-3", "SUBS_r53_058S_Y.png", "원복 후 058S 100011 - 0행"),
    ("5-3", "SUBS_r53_060S_X.png", "원복 후 060S 117786 - 0행"),
]:
    c = ws.cell(row=r, column=1, value=tag)
    c.font = BOLD
    r = image(ws, r, fn, cap, col="B")

for sh in wb.worksheets:
    sh.sheet_view.showGridLines = False

wb.save(OUT)
print("saved:", OUT)
print("sheets:", [s.title for s in wb.worksheets])
print("missing shots:", missing)
